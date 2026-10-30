import io
from pathlib import Path

import pytest
from docx import Document

from app.services.storage import DOCX_CONTENT_TYPE, PDF_CONTENT_TYPE
from app.services.text_extraction import TextExtractionError, extract_text, normalize_whitespace


def _text_pdf(*pages: str) -> bytes:
    """A minimal valid PDF with one line of Helvetica text per page."""
    objects: list[bytes] = []
    page_ids = [3 + 2 * i for i in range(len(pages))]
    font_id = 3 + 2 * len(pages)
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    for pid, text in zip(page_ids, pages, strict=True):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {pid + 1} 0 R >>".encode()
        )
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream))
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n%s\nendobj\n" % (number, body))
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
        % (
            len(objects) + 1,
            xref,
        )
    )
    return out.getvalue()


def _docx(path: Path, *paragraphs: str, table: list[list[str]] | None = None) -> Path:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    document.save(str(path))
    return path


def test_extracts_pdf_text_from_every_page(tmp_path: Path) -> None:
    path = tmp_path / "c.pdf"
    path.write_bytes(_text_pdf("1. Term. One year.", "2. Fees. Net 30."))

    text = extract_text(path, PDF_CONTENT_TYPE)

    assert "1. Term. One year." in text
    assert "2. Fees. Net 30." in text
    assert text.index("Term") < text.index("Fees")


def test_extracts_docx_paragraphs_and_tables(tmp_path: Path) -> None:
    path = _docx(
        tmp_path / "c.docx",
        "Master Agreement",
        "",
        "",
        "1.  Term.\tOne   year.",
        table=[["Fee", "$100"], ["Due", "Net 30"]],
    )

    text = extract_text(path, DOCX_CONTENT_TYPE)

    assert text == "Master Agreement\n\n1. Term. One year.\nFee | $100\nDue | Net 30"


def test_empty_documents_are_rejected(tmp_path: Path) -> None:
    path = _docx(tmp_path / "empty.docx", "", "   ")
    with pytest.raises(TextExtractionError, match="No extractable text"):
        extract_text(path, DOCX_CONTENT_TYPE)


def test_textless_pdf_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    path.write_bytes(_text_pdf(""))
    with pytest.raises(TextExtractionError, match="No extractable text"):
        extract_text(path, PDF_CONTENT_TYPE)


@pytest.mark.parametrize(
    ("content", "content_type"),
    [
        (b"%PDF-1.7 garbage", PDF_CONTENT_TYPE),
        (b"PK\x03\x04 garbage", DOCX_CONTENT_TYPE),
    ],
)
def test_corrupt_files_are_rejected(tmp_path: Path, content: bytes, content_type: str) -> None:
    path = tmp_path / "corrupt"
    path.write_bytes(content)
    with pytest.raises(TextExtractionError):
        extract_text(path, content_type)


def test_unsupported_content_type(tmp_path: Path) -> None:
    with pytest.raises(TextExtractionError, match="No text extractor"):
        extract_text(tmp_path / "x.txt", "text/plain")


def test_normalize_whitespace() -> None:
    assert normalize_whitespace("  a \xa0 b \r\n\r\n\r\n\r\nc\t ") == "a b\n\nc"

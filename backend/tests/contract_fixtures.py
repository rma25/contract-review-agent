"""Synthetic contract fixtures from evals/fixtures, rendered to real DOCX files for upload.

The fixtures are Markdown using only `#`/`##` headings and plain paragraphs.
"""

import io
from pathlib import Path
from typing import Any

import yaml
from docx import Document

from app.core.config import REPO_ROOT

FIXTURES_DIR = REPO_ROOT / "evals" / "fixtures"
HIGH_RISK = "high_risk_vendor.md"
MEDIUM_RISK = "medium_risk_services.md"
LOW_RISK = "low_risk_saas.md"


def render_docx(md_path: Path) -> bytes:
    document = Document()
    paragraph: list[str] = []

    def flush() -> None:
        if paragraph:
            document.add_paragraph(" ".join(paragraph))
            paragraph.clear()

    for line in md_path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if stripped.startswith("#"):
            flush()
            level = len(stripped) - len(stripped.lstrip("#"))
            document.add_heading(stripped[level:].strip(), level=level)
        elif stripped:
            paragraph.append(stripped)
        else:
            flush()
    flush()

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def fixture_docx(name: str) -> bytes:
    return render_docx(FIXTURES_DIR / name)


def expected_labels() -> dict[str, Any]:
    return yaml.safe_load((FIXTURES_DIR / "expected.yaml").read_text(encoding="utf-8"))

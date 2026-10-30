import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.services.storage import DOCX_CONTENT_TYPE, PDF_CONTENT_TYPE

_SPACES = re.compile(r"[ \t\f\v\xa0]+")
_BLANK_LINES = re.compile(r"\n{3,}")


class TextExtractionError(Exception):
    """The file could not be read, or held no extractable text (e.g. a scanned PDF)."""


def extract_text(path: Path, content_type: str) -> str:
    if content_type == PDF_CONTENT_TYPE:
        raw = _pdf_text(path)
    elif content_type == DOCX_CONTENT_TYPE:
        raw = _docx_text(path)
    else:
        raise TextExtractionError(f"No text extractor for {content_type}")
    text = normalize_whitespace(raw)
    if not text:
        raise TextExtractionError("No extractable text found")
    return text


def normalize_whitespace(text: str) -> str:
    lines = (_SPACES.sub(" ", line).strip() for line in text.replace("\r\n", "\n").split("\n"))
    return _BLANK_LINES.sub("\n\n", "\n".join(lines)).strip()


def _pdf_text(path: Path) -> str:
    try:
        reader = PdfReader(path)
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except (PdfReadError, ValueError, KeyError, TypeError) as exc:
        raise TextExtractionError("PDF could not be read") from exc


def _docx_text(path: Path) -> str:
    try:
        document = Document(str(path))
    except Exception as exc:  # python-docx raises assorted zip, XML and key errors
        raise TextExtractionError("DOCX could not be read") from exc
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(blocks)
